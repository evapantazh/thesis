"""
check_all_recordings.py

Scans all recording folders under BASE_DIR, computes quality stats,
and writes a formatted Excel log: recording_log.xlsx

Folder naming convention expected: {subject}_{distance}_{clothing}
Example: andreas_800_tshirt, Sub01_1200_hoodie
"""

import numpy as np
import cv2
import pandas as pd
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── CONFIG ──────────────────────────────────────────────────────────────
BASE_DIR = Path(r"D:\recordings")
OUTPUT_XLS = Path(r"C:\Projects\thesis\data") / "recording_log.xlsx"
MAX_GAP_MS = 100   # threshold to flag a frame drop
# ────────────────────────────────────────────────────────────────────────


def check_recording(folder: Path) -> dict:
    """Run all checks on a single recording folder. Returns a dict of stats."""
    result = {
        "recording":       folder.name,
        "subject":         "",
        "distance_mm":     "",
        "clothing":        "",
        "total_frames":    "",
        "duration_sec":    "",
        "avg_fps":         "",
        "avg_gap_ms":      "",
        "frame_drops":     "",
        "max_gap_ms":      "",
        "depth_shape":     "",
        "color_shape":     "",
        "depth_dtype":     "",
        "depth_min":       "",
        "depth_max":       "",
        "usable":          "Yes",
        "notes":           "",
    }

    # Parse folder name into subject / distance / clothing
    parts = folder.name.split("_")
    if len(parts) >= 3:
        result["subject"]     = parts[0]
        result["distance_mm"] = parts[1]
        result["clothing"]    = parts[2]

    color_dir = folder / "color"
    depth_dir = folder / "depth"
    ts_file   = folder / "timestamps.csv"

    # Check directories exist
    if not color_dir.exists() or not depth_dir.exists():
        result["notes"]  = "Missing color or depth folder"
        result["usable"] = "No"
        return result

    color_files = sorted(list(color_dir.glob("*.jpg")))
    depth_files = sorted(list(depth_dir.glob("*.npy")))

    if len(color_files) == 0 or len(depth_files) == 0:
        result["notes"]  = "Empty color or depth folder"
        result["usable"] = "No"
        return result

    # File count mismatch
    if len(color_files) != len(depth_files):
        result["notes"]  = f"File count mismatch: {len(color_files)} color vs {len(depth_files)} depth"
        result["usable"] = "Check"

    # First frame stats
    try:
        depth = np.load(str(depth_files[0]))
        color = cv2.imread(str(color_files[0]))
        result["depth_shape"] = str(depth.shape)
        result["color_shape"] = str(color.shape) if color is not None else "read error"
        result["depth_dtype"] = str(depth.dtype)
        result["depth_min"]   = int(depth.min())
        result["depth_max"]   = int(depth.max())
    except Exception as e:
        result["notes"] += f" | Frame read error: {e}"

    # Timestamps
    if not ts_file.exists():
        result["notes"] += " | No timestamps.csv"
        return result

    try:
        df = pd.read_csv(ts_file)

        # Support both old (timestamp) and new (depth_timestamp) column names
        ts_col = "depth_timestamp" if "depth_timestamp" in df.columns else "timestamp"

        df["gap"] = df[ts_col].diff()
        total     = len(df)
        avg_fps   = round(1000 / df["gap"].mean(), 1)
        avg_gap   = round(df["gap"].mean(), 1)
        drops     = int((df["gap"] > MAX_GAP_MS).sum())
        max_gap   = round(df["gap"].max(), 1)
        duration  = round((df[ts_col].iloc[-1] - df[ts_col].iloc[0]) / 1000, 1)

        result["total_frames"] = total
        result["duration_sec"] = duration
        result["avg_fps"]      = avg_fps
        result["avg_gap_ms"]   = avg_gap
        result["frame_drops"]  = drops
        result["max_gap_ms"]   = max_gap

        # Auto-flag problems
        issues = []
        if drops > 5:
            issues.append(f"{drops} frame drops")
        if duration < 60:
            issues.append(f"short recording ({duration}s)")
        if avg_fps < 12 or avg_fps > 20:
            issues.append(f"unusual FPS ({avg_fps})")
        if result["depth_max"] == 0:
            issues.append("depth all zeros")

        if issues:
            result["usable"] = "Check"
            result["notes"] += " | ".join(issues)

    except Exception as e:
        result["notes"] += f" | CSV error: {e}"

    return result


def write_excel(rows: list, output_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Recording Log"

    # Colors
    HDR_FILL  = PatternFill("solid", start_color="0D1B2A")
    YES_FILL  = PatternFill("solid", start_color="D5F5E3")
    NO_FILL   = PatternFill("solid", start_color="FADBD8")
    CHK_FILL  = PatternFill("solid", start_color="FEF9E7")
    ALT_FILL  = PatternFill("solid", start_color="F0F4F8")
    thin      = Side(style="thin", color="CCCCCC")
    border    = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = [
        "Recording", "Subject", "Distance (mm)", "Clothing",
        "Total Frames", "Duration (s)", "Avg FPS", "Avg Gap (ms)",
        "Frame Drops", "Max Gap (ms)", "Depth Shape", "Color Shape",
        "Depth Dtype", "Depth Min", "Depth Max", "Usable", "Notes"
    ]

    keys = [
        "recording", "subject", "distance_mm", "clothing",
        "total_frames", "duration_sec", "avg_fps", "avg_gap_ms",
        "frame_drops", "max_gap_ms", "depth_shape", "color_shape",
        "depth_dtype", "depth_min", "depth_max", "usable", "notes"
    ]

    col_widths = [28, 12, 14, 10, 13, 12, 9, 13, 13, 13, 16, 18, 12, 10, 10, 8, 40]

    # Header row
    for col, (hdr, w) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=1, column=col, value=hdr)
        cell.font      = Font(bold=True, color="FFFFFF", name="Arial", size=10)
        cell.fill      = HDR_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border    = border
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.row_dimensions[1].height = 30

    # Data rows
    for row_idx, rec in enumerate(rows, 2):
        fill = ALT_FILL if row_idx % 2 == 0 else PatternFill("solid", start_color="FFFFFF")
        for col, key in enumerate(keys, 1):
            val  = rec.get(key, "")
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.font      = Font(name="Arial", size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border    = border
            cell.fill      = fill

        # Color the Usable column
        usable_cell = ws.cell(row=row_idx, column=16)
        if rec.get("usable") == "Yes":
            usable_cell.fill = YES_FILL
            usable_cell.font = Font(name="Arial", size=10, bold=True, color="1E8449")
        elif rec.get("usable") == "No":
            usable_cell.fill = NO_FILL
            usable_cell.font = Font(name="Arial", size=10, bold=True, color="C0392B")
        else:
            usable_cell.fill = CHK_FILL
            usable_cell.font = Font(name="Arial", size=10, bold=True, color="D35400")

        ws.row_dimensions[row_idx].height = 20

    # Summary row at the bottom
    last = len(rows) + 2
    ws.cell(row=last, column=1, value="TOTAL").font = Font(bold=True, name="Arial")
    ws.cell(row=last, column=5, value=f"=SUM(E2:E{last-1})").font = Font(bold=True, name="Arial")
    ws.cell(row=last, column=9, value=f"=SUM(I2:I{last-1})").font = Font(bold=True, name="Arial")

    # Freeze header row
    ws.freeze_panes = "A2"

    wb.save(output_path)
    print(f"Saved: {output_path}")


def main():
    print(f"Scanning: {BASE_DIR}")

    # Find all folders that contain a 'depth' subfolder
    recording_folders = sorted([
        f for f in BASE_DIR.iterdir()
        if f.is_dir() 
        and (f / "depth").exists()
        and not f.name.startswith("Camera_")
    ])
    if not recording_folders:
        print("No recording folders found. Check BASE_DIR path.")
        return

    print(f"Found {len(recording_folders)} recording folders\n")

    rows = []
    for folder in recording_folders:
        print(f"  Checking: {folder.name} ...", end=" ")
        rec = check_recording(folder)
        rows.append(rec)
        status = f"{rec['total_frames']} frames | {rec['avg_fps']} fps | {rec['frame_drops']} drops | {rec['usable']}"
        print(status)

    write_excel(rows, OUTPUT_XLS)
    print(f"\nDone. {len(rows)} recordings logged.")


if __name__ == "__main__":
    main()